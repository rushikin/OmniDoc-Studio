"""
Word Document Generator Module
Takes a list of AI-enhanced page elements and reconstructs them into a
properly formatted .docx file with native Word tables.
"""

import os
from typing import List, Dict, Any

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ─── Helpers ──────────────────────────────────────────────────────────────────

def set_cell_background(cell, hex_color: str) -> None:
    """Set a table cell's background fill color."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_table_border(table) -> None:
    """Apply visible borders to all cells of a Word table."""
    tbl = table._tbl
    # Safely get or create tblPr element
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)

    # Remove any existing border definition to avoid duplicates
    existing = tblPr.find(qn("w:tblBorders"))
    if existing is not None:
        tblPr.remove(existing)

    tblBorders = OxmlElement("w:tblBorders")
    for border_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{border_name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "4472C4")
        tblBorders.append(border)
    tblPr.append(tblBorders)


# ─── Main Generator Class ─────────────────────────────────────────────────────

class WordGenerator:
    """
    Builds a polished .docx file from the list of enhanced page elements.
    """

    def __init__(self):
        self.doc = Document()
        self._setup_document_styles()

    def _setup_document_styles(self) -> None:
        """Configure global document margins, default font, and page-number footer."""
        section = self.doc.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)

        # Default body font
        style = self.doc.styles["Normal"]
        font = style.font
        font.name = "Calibri"
        font.size = Pt(11)

        # Add centered page-number footer
        self._add_page_number_footer(section)

    def _add_page_number_footer(self, section) -> None:
        """Insert a centered 'Page X of Y' footer using Word field codes."""
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.clear()

        def _add_field(run_elem, field_type: str):
            fldChar_begin = OxmlElement("w:fldChar")
            fldChar_begin.set(qn("w:fldCharType"), "begin")
            instrText = OxmlElement("w:instrText")
            instrText.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            instrText.text = f" {field_type} "
            fldChar_end = OxmlElement("w:fldChar")
            fldChar_end.set(qn("w:fldCharType"), "end")
            run_elem.append(fldChar_begin)
            run_elem.append(instrText)
            run_elem.append(fldChar_end)

        r1 = OxmlElement("w:r")
        t1 = OxmlElement("w:t")
        t1.text = "Page "
        t1.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r1.append(t1)
        p._p.append(r1)

        r_page = OxmlElement("w:r")
        _add_field(r_page, "PAGE")
        p._p.append(r_page)

        r2 = OxmlElement("w:r")
        t2 = OxmlElement("w:t")
        t2.text = " of "
        t2.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        r2.append(t2)
        p._p.append(r2)

        r_numpages = OxmlElement("w:r")
        _add_field(r_numpages, "NUMPAGES")
        p._p.append(r_numpages)

    def generate_toc(self, pages_elements: List[List[Dict[str, Any]]]) -> None:
        """
        Generate a Table of Contents from title elements found across all pages.
        Call this BEFORE add_page calls to prepend the TOC at document start.

        Args:
            pages_elements: All page elements (list of lists) to scan for titles.
        """
        titles = []
        for page_idx, elements in enumerate(pages_elements):
            for el in elements:
                if el.get("type") == "title" and el.get("content", "").strip():
                    titles.append((page_idx + 1, el["content"].strip()))

        if len(titles) < 2:
            return  # Not worth a TOC for fewer than 2 titles

        toc_heading = self.doc.add_paragraph("Table of Contents", style="Heading 1")
        toc_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for page_num, title_text in titles:
            p = self.doc.add_paragraph(style="Normal")
            run = p.add_run(f"  {title_text}")
            run.font.size = Pt(11)
            tab_run = p.add_run(f"\t{page_num}")
            tab_run.font.size = Pt(11)
            tab_run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
        self.doc.add_paragraph()  # Spacer
        self.doc.add_page_break()

    def add_page(self, page_index: int, elements: List[Dict[str, Any]]) -> None:
        """
        Write all elements from a single page into the document.
        Adds a page break between pages (except after the last page).

        Args:
            page_index: 0-based page number.
            elements: List of AI-enhanced elements for this page.
        """
        if page_index > 0:
            self.doc.add_page_break()

        print(f"[Word Generator] Writing page {page_index + 1} ({len(elements)} element(s))...")

        for el in elements:
            el_type = el.get("type", "text")
            content = el.get("content", "")

            if el_type == "title":
                self._add_title(content)
            elif el_type == "header":
                self._add_heading(content, level=2)
            elif el_type == "table":
                self._add_table(content)
            elif el_type in ("text", "footer"):
                self._add_paragraph(content)

    def _add_title(self, text: str) -> None:
        """Add a styled document title paragraph."""
        p = self.doc.add_heading(text, level=1)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.runs[0] if p.runs else p.add_run(text)
        run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)  # Dark blue

    def _add_heading(self, text: str, level: int = 2) -> None:
        """Add a section heading."""
        p = self.doc.add_heading(text, level=level)
        run = p.runs[0] if p.runs else p.add_run(text)
        run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)  # Medium blue

    def _add_paragraph(self, text: str) -> None:
        """Add a normal text paragraph."""
        if not text.strip():
            return
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(11)

    def _add_table(self, grid: Any) -> None:
        """
        Add a Word table from a 2D list of strings.
        The first row is treated as a header row with colored background.

        Args:
            grid: List[List[str]] — the table data.
        """
        if not grid or not isinstance(grid, list):
            self._add_paragraph("[Table data unavailable]")
            return

        # Ensure every row is a list
        if not isinstance(grid[0], list):
            grid = [[str(cell) for cell in row] for row in grid]

        num_rows = len(grid)
        num_cols = max(len(row) for row in grid) if grid else 1

        # Pad shorter rows to uniform width
        for i, row in enumerate(grid):
            if len(row) < num_cols:
                grid[i] = row + [""] * (num_cols - len(row))

        table = self.doc.add_table(rows=num_rows, cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_border(table)

        for row_idx, row_data in enumerate(grid):
            row = table.rows[row_idx]
            for col_idx, cell_text in enumerate(row_data):
                cell = row.cells[col_idx]
                cell.text = str(cell_text)

                # Style header row
                if row_idx == 0:
                    set_cell_background(cell, "2E74B5")
                    for run in cell.paragraphs[0].runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.font.size = Pt(10)
                else:
                    # Alternate row shading
                    bg = "DEEAF1" if row_idx % 2 == 0 else "FFFFFF"
                    set_cell_background(cell, bg)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(10)

        # Add spacing after table
        self.doc.add_paragraph()

    def save(self, output_path: str) -> None:
        """Save the document to disk."""
        os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
        self.doc.save(output_path)
        print(f"[Word Generator] Saved: {output_path}")
